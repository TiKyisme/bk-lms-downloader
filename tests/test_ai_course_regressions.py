import json
import subprocess
import sys
from dataclasses import asdict
from pathlib import Path
import zipfile

from docx import Document
from openpyxl import Workbook
import pytest

from bklms_downloader.ai_prepare import AIBatchPreparer, AICoursePreparer, AIPreparationError
from bklms_downloader.ai_prepare import AIBatchPreparationResult, AICoursePreparationResult
from bklms_downloader.ai_study_pack import (
    NAVIGATION_FILES,
    calculate_study_pack_source_metrics,
    validate_ai_study_pack,
    write_study_navigation,
)
from bklms_downloader.course_roots import resolve_course_root
from bklms_downloader.coursewave import CoursewaveCatalogResult
from bklms_downloader.gui import App
from bklms_downloader.models import Course
from bklms_downloader.study_pack_refresh import (
    PROCESSING_FINGERPRINT,
    CoursewaveEnricher,
    plan_refresh,
    snapshot_sources,
)
from bklms_downloader.ai_sources import inventory_course_sources


def _write_docx(path: Path, marker: str) -> None:
    document = Document()
    document.add_heading("Course teaching notes", level=1)
    document.add_paragraph(f"{marker} explains the core food quality management process.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Control point"
    table.rows[0].cells[1].text = marker
    document.save(path)


def _write_xlsx(path: Path, marker: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quality controls"
    sheet.append(["Control", "Value"])
    sheet.append([marker, "12 percent"])
    sheet.append(["Formula", "=SUM(B2:B2)"])
    workbook.save(path)


def _pack_members(pack: Path) -> tuple[list[dict], list[dict], dict, dict]:
    with zipfile.ZipFile(pack) as archive:
        records = [json.loads(line) for line in archive.read("meta/documents.jsonl").decode().splitlines() if line]
        chunks = [json.loads(line) for line in archive.read("meta/corpus.jsonl").decode().splitlines() if line]
        manifest = json.loads(archive.read("meta/pack_manifest.json"))
        exam_manifest = json.loads(archive.read("meta/exam_manifest.json"))
    return records, chunks, manifest, exam_manifest


def _write_empty_owned_pack(
    pack: Path,
    source_root: Path,
    course_name: str = "Course",
    *,
    include_coursewave_exam: bool = False,
) -> None:
    workspace = pack.parent / f"{pack.stem}-control"
    workspace.mkdir()
    (workspace / "meta").mkdir()
    exam_entries = []
    if include_coursewave_exam:
        exam_path = workspace / "sources" / "past_exams" / "final-only__Final.pdf"
        exam_path.parent.mkdir(parents=True)
        exam_path.write_bytes(b"synthetic exam evidence")
        exam_entries = [{
            "stable_id": "final-only",
            "original_name": "Final.pdf",
            "source_role": "past_exam",
            "retained_source_path": "sources/past_exams/final-only__Final.pdf",
        }]
    write_study_navigation(workspace, course_name, [], [], exams=exam_entries)
    for name in ("documents.jsonl", "corpus.jsonl"):
        (workspace / "meta" / name).write_text("", encoding="utf-8")
    (workspace / "meta" / "visual_manifest.json").write_text("[]\n", encoding="utf-8")
    (workspace / "meta" / "source_inventory.json").write_text(
        json.dumps(inventory_course_sources(source_root).to_dict()), encoding="utf-8"
    )
    (workspace / "meta" / "stats.json").write_text(
        json.dumps({"source_root": ".", "source_paths_relative": True}), encoding="utf-8"
    )
    (workspace / "meta" / "pack_manifest.json").write_text(
        json.dumps({
            "processing_fingerprint": PROCESSING_FINGERPRINT,
            "course_name": course_name,
            "one_course_only": True,
        }),
        encoding="utf-8",
    )
    source_snapshots = [asdict(item) for item in snapshot_sources(source_root).values()]
    (workspace / "meta" / "source_manifest.json").write_text(
        json.dumps({"sources": source_snapshots}), encoding="utf-8"
    )
    (workspace / "meta" / "exam_manifest.json").write_text(
        json.dumps({"exams": exam_entries, "diagnostics": {"stage": "no_match"}}), encoding="utf-8"
    )
    with zipfile.ZipFile(pack, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for path in workspace.rglob("*"):
            if path.is_file():
                archive.write(path, path.relative_to(workspace).as_posix())


def test_docx_only_course_produces_traceable_documents_and_chunks(tmp_path: Path):
    course_root = tmp_path / "DEMO9001"
    course_root.mkdir()
    _write_docx(course_root / "lecture.docx", "DOCX_ONLY_MARKER")

    pack = AICoursePreparer().prepare(course_root, course_code="DEMO9001")

    records, chunks, manifest, _ = _pack_members(pack)
    office = [record for record in records if record["source_type"] == "word_document"]
    assert len(office) == 1 and office[0]["status"] == "ready"
    assert office[0]["source_copy_path"]
    assert chunks and any("DOCX_ONLY_MARKER" in chunk["text"] for chunk in chunks)
    assert manifest["source_metrics"]["ready_source_count"] >= 1
    assert manifest["source_metrics"]["chunk_count"] >= 1
    assert manifest["source_inventory"]["extension_counts"][".docx"] == 1
    with zipfile.ZipFile(pack) as archive:
        assert "meta/source_inventory.json" in archive.namelist()


def test_xlsx_only_course_preserves_values_formulas_and_original(tmp_path: Path):
    course_root = tmp_path / "DEMO9002"
    course_root.mkdir()
    _write_xlsx(course_root / "quality.xlsx", "XLSX_ONLY_MARKER")

    pack = AICoursePreparer().prepare(course_root, course_code="DEMO9002")

    records, chunks, manifest, _ = _pack_members(pack)
    spreadsheet = [record for record in records if record["source_type"] == "spreadsheet"]
    assert len(spreadsheet) == 1 and spreadsheet[0]["status"] == "ready"
    assert spreadsheet[0]["source_copy_path"]
    assert chunks and any("=SUM(B2:B2)" in chunk["text"] for chunk in chunks)
    assert any("XLSX_ONLY_MARKER" in chunk["text"] for chunk in chunks)
    assert manifest["source_metrics"]["meaningful_lecturer_source_count"] >= 1


def test_coursewave_no_match_is_warning_and_local_material_still_succeeds(tmp_path: Path):
    course_root = tmp_path / "Demo course local folder"
    course_root.mkdir()
    _write_docx(course_root / "lecture.docx", "LOCAL_BK_LMS_CONTENT")

    class NoMatchCatalog:
        def fetch_catalog_result(self):
            return CoursewaveCatalogResult("ok", courses=())

    class CoursewavePreparer(AICoursePreparer):
        def prepare(self, *args, **kwargs):
            kwargs["coursewave_enricher"] = CoursewaveEnricher(catalog_client=NoMatchCatalog())
            return super().prepare(*args, **kwargs)

    course = Course("demo9002", "https://lms.hcmut.edu.vn/course/view.php?id=9002", str(course_root), "Demo course", "DEMO9002")
    batch = AIBatchPreparer(lambda: CoursewavePreparer()).prepare_courses(
        [course], lambda _course: course_root, enrich_exams=True
    )

    assert len(batch.succeeded) == 1 and batch.failed == []
    result = batch.succeeded[0]
    assert result.succeeded_with_warnings
    assert any("BK-LMS" in warning for warning in result.warnings)
    records, chunks, manifest, exam_manifest = _pack_members(result.output)
    assert any(record["source_type"] == "word_document" and record["status"] == "ready" for record in records)
    assert chunks
    assert manifest["coursewave"]["stage"] == "no_match"
    assert exam_manifest["exams"] == []


def test_unsupported_only_and_empty_courses_fail_with_actionable_inventory(tmp_path: Path):
    unsupported = tmp_path / "unsupported"
    unsupported.mkdir()
    (unsupported / "legacy.doc").write_bytes(b"legacy Word")
    (unsupported / "legacy.xls").write_bytes(b"legacy Excel")
    preparer = AICoursePreparer()

    with pytest.raises(AIPreparationError, match=r"\.doc: 1.*\.xls: 1"):
        preparer.prepare(unsupported)
    assert preparer.last_source_inventory["unsupported_files"] == 2
    assert not list(tmp_path.glob("*_AI_Study_Pack*.zip"))

    empty = tmp_path / "empty"
    empty.mkdir()
    with pytest.raises(AIPreparationError, match="Study Pack"):
        preparer.prepare(empty)


def test_corrupt_pdf_is_warning_when_other_teaching_source_is_usable(tmp_path: Path):
    course_root = tmp_path / "mixed"
    course_root.mkdir()
    (course_root / "good.txt").write_text(
        "Lecturer notes explain the process control requirements and the measurement method.",
        encoding="utf-8",
    )
    (course_root / "corrupt.pdf").write_bytes(b"not a PDF package")

    preparer = AICoursePreparer()
    pack = preparer.prepare(course_root, course_code="DEMO9002")

    records, chunks, manifest, _ = _pack_members(pack)
    assert chunks
    assert any(record["status"] == "error" and record["source_path"] == "corrupt.pdf" for record in records)
    assert manifest["source_metrics"]["source_error_count"] == 1
    assert preparer.last_source_warnings
    assert str(tmp_path) not in json.dumps(manifest)


def test_old_semantically_empty_pack_is_invalid_and_rebuilt_atomically(tmp_path: Path):
    course_root = tmp_path / "course"
    course_root.mkdir()
    (course_root / "notes.txt").write_text(
        "The lecturer explains the main database normalization rules with examples.",
        encoding="utf-8",
    )
    old_pack = tmp_path / "Course_AI_Study_Pack.zip"
    _write_empty_owned_pack(old_pack, course_root, "Course")
    before = old_pack.read_bytes()

    assert plan_refresh(old_pack, course_root).state == "invalid"
    with zipfile.ZipFile(old_pack) as archive:
        old_workspace = tmp_path / "old-unpacked"
        old_workspace.mkdir()
        archive.extractall(old_workspace)
    old_report = validate_ai_study_pack(old_workspace)
    assert "No usable course teaching sources were included." in old_report.errors

    rebuilt = AICoursePreparer().prepare(course_root, existing_pack=old_pack, course_name="Course")

    assert rebuilt == old_pack
    assert old_pack.read_bytes() != before
    assert plan_refresh(old_pack, course_root).state == "up_to_date"
    records, chunks, manifest, _ = _pack_members(old_pack)
    assert records and chunks
    assert manifest["source_metrics"]["viable"] is True


def test_old_processing_fingerprint_forces_rebuild(tmp_path: Path):
    from bklms_downloader.study_pack_refresh import PROCESSING_FINGERPRINT

    course_root = tmp_path / "course"
    course_root.mkdir()
    (course_root / "notes.txt").write_text(
        "The lecturer explains the main database normalization rules with examples.",
        encoding="utf-8",
    )
    pack = AICoursePreparer().prepare(course_root, course_name="Course")
    temp_pack = tmp_path / "old-fingerprint.zip"
    with zipfile.ZipFile(pack) as source, zipfile.ZipFile(temp_pack, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "meta/pack_manifest.json":
                manifest = json.loads(data)
                manifest["processing_fingerprint"] = "bklms-study-pack-v1.3"
                data = json.dumps(manifest).encode("utf-8")
            target.writestr(name, data)
    temp_pack.replace(pack)

    plan = plan_refresh(pack, course_root)
    assert plan.state == "legacy" and plan.requires_rebuild

    AICoursePreparer().prepare(course_root, existing_pack=pack, course_name="Course")

    with zipfile.ZipFile(pack) as archive:
        manifest = json.loads(archive.read("meta/pack_manifest.json"))
    assert manifest["processing_fingerprint"] == PROCESSING_FINGERPRINT
    assert manifest["source_metrics"]["viable"] is True


def test_coursewave_exams_alone_do_not_satisfy_pack_viability(tmp_path: Path):
    course_root = tmp_path / "course"
    course_root.mkdir()
    pack = tmp_path / "exam-only.zip"
    _write_empty_owned_pack(pack, course_root, include_coursewave_exam=True)
    workspace = tmp_path / "exam-only-workspace"
    workspace.mkdir()
    with zipfile.ZipFile(pack) as archive:
        archive.extractall(workspace)

    report = validate_ai_study_pack(workspace)

    assert "No usable course teaching sources were included." in report.errors
    assert report.metrics["ready_source_count"] == 0
    assert report.metrics["chunk_count"] == 0
    assert report.metrics["retained_lecturer_source_count"] == 0


def test_failed_semantic_refresh_keeps_previous_pack_bytes_and_reference(tmp_path: Path):
    course_root = tmp_path / "course"
    course_root.mkdir()
    (course_root / "notes.txt").write_text("Good teaching content describes concepts and examples.", encoding="utf-8")
    pack = AICoursePreparer().prepare(course_root, course_name="Course")
    previous = pack.read_bytes()
    (course_root / "notes.txt").unlink()
    (course_root / "bad.docx").write_bytes(b"broken document")

    with pytest.raises(AIPreparationError):
        AICoursePreparer().prepare(course_root, course_name="Course", existing_pack=pack)

    assert pack.read_bytes() == previous
    assert plan_refresh(pack, course_root).state == "dirty"


def test_cross_course_batch_resolves_metadata_roots_and_keeps_packs_isolated(tmp_path: Path):
    shared = tmp_path / "BK_LMS_Data"
    first_root = shared / "renamed-download-one"
    second_root = shared / "renamed-download-two"
    for root, url, title, marker in (
        (first_root, "https://lms.hcmut.edu.vn/course/view.php?id=9001", "DEMO9001 Course A", "ONLY_DEMO_A_CONTENT"),
        (second_root, "https://lms.hcmut.edu.vn/course/view.php?id=9002", "DEMO9002 Course B", "ONLY_DEMO_B_CONTENT"),
    ):
        root.mkdir(parents=True)
        meta = root / "_meta"
        meta.mkdir()
        (meta / "course_structure.json").write_text(
            json.dumps({"root_course": title, "root_course_url": url, "courses": []}),
            encoding="utf-8",
        )
        (root / "notes.txt").write_text(
            f"{marker}: the lecturer explains the course process and its important concepts.",
            encoding="utf-8",
        )

    first = Course("c1", "https://lms.hcmut.edu.vn/course/view.php?id=9001", str(shared), "Course 9001", "DEMO9001")
    second = Course("c2", "https://lms.hcmut.edu.vn/course/view.php?id=9002", str(shared), "Course 9002", "DEMO9002")
    batch = AIBatchPreparer().prepare_courses([first, second], resolve_course_root)

    assert len(batch.succeeded) == 2 and batch.failed == []
    for result, own_marker, other_marker in (
        (batch.succeeded[0], "ONLY_DEMO_A_CONTENT", "ONLY_DEMO_B_CONTENT"),
        (batch.succeeded[1], "ONLY_DEMO_B_CONTENT", "ONLY_DEMO_A_CONTENT"),
    ):
        with zipfile.ZipFile(result.output) as archive:
            content = "\n".join(
                archive.read(name).decode("utf-8", errors="replace")
                for name in archive.namelist()
                if name.endswith((".md", ".jsonl", ".json"))
            )
        assert own_marker in content
        assert other_marker not in content


def test_packaged_diagnostic_cli_outputs_only_relative_safe_metrics(tmp_path: Path):
    course_root = tmp_path / "diagnostic-course"
    course_root.mkdir()
    _write_docx(course_root / "lecture.docx", "DO_NOT_PRINT_SOURCE_CONTENT")
    (course_root / "token=private-session.bin").write_bytes(b"unsupported diagnostic fixture")

    completed = subprocess.run(
        [sys.executable, "app.py", "--diagnose-ai-course", str(course_root)],
        check=False,
        capture_output=True,
        text=True,
        timeout=60,
    )

    assert completed.returncode == 0
    report = json.loads(completed.stdout)
    assert report["resolved_root_status"] == "resolved_legacy_explicit_folder"
    assert report["source_inventory"]["extension_counts"][".docx"] == 1
    assert report["ready_source_count"] > 0
    assert report["chunk_count"] > 0
    assert report["failed_relative_files"] == []
    assert str(course_root) not in completed.stdout
    assert "DO_NOT_PRINT_SOURCE_CONTENT" not in completed.stdout
    assert "private-session" not in completed.stdout
    assert "token=[redacted]" in completed.stdout


def test_packaged_diagnostic_refuses_shared_parent_with_multiple_course_metadata(tmp_path: Path):
    shared = tmp_path / "shared"
    for index in (1, 2):
        root = shared / f"folder-{index}"
        (root / "_meta").mkdir(parents=True)
        (root / "_meta" / "course_structure.json").write_text(
            json.dumps({"root_course_url": f"https://lms.hcmut.edu.vn/course/view.php?id={index}"}),
            encoding="utf-8",
        )
        (root / "notes.txt").write_text("different course teaching notes", encoding="utf-8")

    completed = subprocess.run(
        [sys.executable, "app.py", "--diagnose-ai-course", str(shared)],
        check=False,
        capture_output=True,
        text=True,
        timeout=60,
    )

    assert completed.returncode == 1
    report = json.loads(completed.stdout)
    assert report["resolved_root_status"] == "ambiguous_multiple_courses"
    assert "failure_reason" in report
    assert str(shared) not in completed.stdout


def test_diagnostic_cli_writes_safe_report_for_windowed_invocation(tmp_path: Path):
    course_root = tmp_path / "diagnostic-course"
    course_root.mkdir()
    _write_docx(course_root / "lecture.docx", "REPORT_SOURCE_SECRET")
    report_path = tmp_path / "safe-diagnostic.json"

    completed = subprocess.run(
        [
            sys.executable,
            "app.py",
            "--diagnose-ai-course",
            str(course_root),
            "--diagnose-ai-course-output",
            str(report_path),
        ],
        check=False,
        capture_output=True,
        text=True,
        timeout=60,
    )

    assert completed.returncode == 0
    assert completed.stdout == ""
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert report["meaningful_lecturer_source_count"] > 0
    rendered = json.dumps(report)
    assert str(course_root) not in rendered
    assert "REPORT_SOURCE_SECRET" not in rendered


class _FakeVar:
    def set(self, value):
        self.value = value


def _bare_ai_app(store, monkeypatch):
    app = App.__new__(App)
    app.syncing = False
    app.current_course_var = _FakeVar()
    app._progress_total_courses = 1
    app._progress_completed_courses = 0
    app._set_batch_progress_target = lambda *_args, **_kwargs: None
    app._finish_sync_activity = lambda: None
    app._set_busy = lambda *_args, **_kwargs: None
    app._set_summary_message = lambda message: setattr(app, "summary", message)
    app._refresh_course_row = lambda _course_id: None
    app._format_pack_size = lambda _path: ""
    app._log = lambda message: None
    app.store = store
    monkeypatch.setattr("bklms_downloader.gui.messagebox.showinfo", lambda *_args, **_kwargs: None)
    monkeypatch.setattr("bklms_downloader.gui.messagebox.showwarning", lambda *_args, **_kwargs: None)
    monkeypatch.setattr("bklms_downloader.gui.messagebox.showerror", lambda *_args, **_kwargs: None)
    return app


def test_failed_refresh_keeps_existing_pack_reference_and_marks_it_dirty(tmp_path: Path, monkeypatch):
    from bklms_downloader.course_store import CourseStore

    store = CourseStore(tmp_path / "courses.json")
    course = store.add("https://lms.hcmut.edu.vn/course/view.php?id=9002", tmp_path / "out", name="DEMO9002")
    old_pack = tmp_path / "old-good-pack.zip"
    old_pack.write_bytes(b"previous pack")
    store.update_study_pack(course.id, path=old_pack, status="up_to_date")
    course = store.get(course.id)
    result = AICoursePreparationResult(course, error="No usable teaching sources")
    app = _bare_ai_app(store, monkeypatch)

    App._complete_ai_batch(app, AIBatchPreparationResult([result]))

    updated = CourseStore(store.path).get(course.id)
    assert updated.study_pack_path == str(old_pack)
    assert updated.study_pack_status == "dirty"
    assert App._status_text(updated)[0] == "AI c\u1ea7n c\u1eadp nh\u1eadt"


def test_success_with_coursewave_warning_is_counted_as_success_and_warning(tmp_path: Path, monkeypatch):
    from bklms_downloader.course_store import CourseStore

    store = CourseStore(tmp_path / "courses.json")
    course = store.add("https://lms.hcmut.edu.vn/course/view.php?id=9002", tmp_path / "out", name="DEMO9002")
    pack = tmp_path / "DEMO9002_AI_Study_Pack.zip"
    pack.write_bytes(b"synthetic successful pack")
    result = AICoursePreparationResult(course, output=pack, warnings=("Coursewave no match",))
    app = _bare_ai_app(store, monkeypatch)

    App._complete_ai_batch(app, AIBatchPreparationResult([result]))

    assert store.get(course.id).study_pack_status == "up_to_date"
    assert "1 course có cảnh báo" in app.summary
    assert "0 course" in app.summary


def test_gui_reports_coursewave_and_source_warnings_separately(tmp_path: Path):
    course = Course("demo9002", "https://lms.hcmut.edu.vn/course/view.php?id=9002", str(tmp_path), "Demo course", "DEMO9002")
    result = AICoursePreparationResult(
        course,
        output=tmp_path / "DEMO9002_AI_Study_Pack.zip",
        warnings=("No Coursewave match; BK-LMS material remains in use.",),
        source_warnings=("One unsupported source was skipped.",),
    )
    app = App.__new__(App)
    app._progress_total_courses = 1
    app._progress_completed_courses = 0
    app._set_batch_progress_target = lambda *_args, **_kwargs: None
    messages = []
    app._log = messages.append
    app._format_pack_size = lambda _path: ""

    App._handle_event(app, {"event": "ai_prepare_course_complete", "result": result, "index": 1, "total": 1})

    assert any(message.startswith("[COURSEWAVE][SKIP]") for message in messages)
    assert any(message.startswith("[AI][WARN]") for message in messages)
