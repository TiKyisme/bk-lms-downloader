import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
from prepare_ai_course import extract_csv, extract_docx, extract_spreadsheet


def test_docx_extraction_preserves_headings_paragraphs_and_table_locators(tmp_path: Path):
    source = tmp_path / "lecture.docx"
    document = Document()
    document.add_heading("Food quality", level=1)
    document.add_paragraph("Quality control compares measured characteristics against requirements.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Measure"
    table.rows[0].cells[1].text = "Target"
    row = table.add_row().cells
    row[0].text = "Moisture"
    row[1].text = "12 percent"
    document.save(source)

    body, units, has_visual = extract_docx(source)

    assert "## Food quality" in body
    assert "Quality control compares" in body
    assert "Moisture" in body and "12 percent" in body
    assert any(locator.startswith("table 1 row 2") for locator, _ in units)
    assert has_visual is False


def test_xlsx_extraction_keeps_sheet_rows_and_formula_text_without_evaluation(tmp_path: Path):
    source = tmp_path / "quality.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Quality measures"
    sheet.append(["Measure", "Value"])
    sheet.append(["Moisture", 12])
    sheet.append(["Calculated total", "=SUM(B2:B2)"])
    workbook.save(source)

    body, units = extract_spreadsheet(source)

    assert "Sheet: Quality measures" in body
    assert "A2: Moisture" in body and "B2: 12" in body
    assert "=SUM(B2:B2)" in body
    assert any(locator == "sheet Quality measures / row 3" for locator, _ in units)


def test_xlsm_is_read_as_data_without_executing_macros(tmp_path: Path):
    source = tmp_path / "macro-enabled.xlsm"
    workbook = Workbook()
    workbook.active["A1"] = "XLSM_ONLY_MARKER"
    workbook.save(source)

    body, units = extract_spreadsheet(source)

    assert "XLSM_ONLY_MARKER" in body
    assert units and units[0][0].startswith("sheet Sheet / row")


def test_csv_extraction_handles_bom_delimiter_and_row_locators(tmp_path: Path):
    source = tmp_path / "measurements.csv"
    source.write_text("Mã;Giá trị\nĐộ ẩm;12\n", encoding="utf-8-sig", newline="")

    body, units = extract_csv(source)

    assert "Mã" in body and "Độ ẩm" in body and "12" in body
    assert units[0][0] == "row 1"
    assert units[1][0] == "row 2"


def test_csv_extraction_falls_back_to_vietnamese_legacy_encoding(tmp_path: Path):
    source = tmp_path / "legacy.csv"
    source.write_bytes(b"Topic;Value\nSugar;\xe1\n")

    body, _units = extract_csv(source)

    assert "\u00e1" in body


def test_empty_office_documents_do_not_create_teaching_text(tmp_path: Path):
    docx_path = tmp_path / "blank.docx"
    Document().save(docx_path)
    xlsx_path = tmp_path / "blank.xlsx"
    Workbook().save(xlsx_path)

    docx_body, docx_units, _ = extract_docx(docx_path)
    xlsx_body, xlsx_units = extract_spreadsheet(xlsx_path)

    assert not docx_body.strip() and docx_units == []
    assert not xlsx_body.strip() and xlsx_units == []
